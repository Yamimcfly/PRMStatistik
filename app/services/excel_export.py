"""
Excel Export mit professioneller Formatierung
"""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime


def apply_header_style(worksheet, row_num=1):
    """Wendet Kopfzeilen-Styling an"""
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    for cell in worksheet[row_num]:
        if cell.value:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment


def apply_column_widths(worksheet, df):
    """Passt Spaltenbreiten automatisch an"""
    for idx, col in enumerate(df.columns, 1):
        max_length = len(str(col))
        for cell in worksheet[get_column_letter(idx)]:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        worksheet.column_dimensions[get_column_letter(idx)].width = adjusted_width


def apply_data_formatting(worksheet, start_row=2):
    """Formatiert Datenzeilen mit Rahmen und Ausrichtung"""
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for row in worksheet.iter_rows(min_row=start_row):
        for cell in row:
            cell.border = thin_border
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")


def add_title_rows(worksheet, title, date_range=None, station="CGN", airline="All Airlines"):
    """Fügt Titelzeilen wie in der Vorlage hinzu"""
    # Titel in Zeile 1
    worksheet.insert_rows(1, 6)
    
    worksheet.merge_cells('A1:H1')
    title_cell = worksheet['A1']
    title_cell.value = title
    title_cell.font = Font(bold=True, size=14, color="366092")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Date Range in Zeile 2
    if date_range:
        worksheet.merge_cells('A2:H2')
        date_cell = worksheet['A2']
        date_cell.value = f"Date Range: {date_range}"
        date_cell.font = Font(size=10)
        date_cell.alignment = Alignment(horizontal="center")
    
    # Airline in Zeile 3
    worksheet.merge_cells('A3:H3')
    airline_cell = worksheet['A3']
    airline_cell.value = f"Airline: {airline}"
    airline_cell.font = Font(size=10)
    airline_cell.alignment = Alignment(horizontal="center")
    
    # Station in Zeile 4
    worksheet['A4'].value = f"Station: {station}"
    worksheet['A4'].font = Font(size=10)
    
    # Version in Zeile 5
    worksheet['A5'].value = f"Version: {datetime.now().strftime('%Y%m%d %H:%M')}"
    worksheet['A5'].font = Font(size=9, color="808080")


def export_with_formatting(writer, df, sheet_name, title=None, date_range=None):
    """
    Exportiert DataFrame mit professioneller Formatierung
    
    Args:
        writer: pd.ExcelWriter Objekt
        df: DataFrame zum Exportieren
        sheet_name: Name des Sheets
        title: Titel für das Sheet (optional)
        date_range: Datumsbereich als String (optional)
    """
    # DataFrame in Excel schreiben
    df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=0)
    worksheet = writer.sheets[sheet_name]
    
    # Wenn Titel vorhanden, füge Titelzeilen hinzu
    if title:
        add_title_rows(worksheet, title, date_range)
        # Kopfzeile ist jetzt in Zeile 7
        apply_header_style(worksheet, row_num=7)
        apply_data_formatting(worksheet, start_row=8)
    else:
        # Standard-Formatierung ohne Titelzeilen
        apply_header_style(worksheet, row_num=1)
        apply_data_formatting(worksheet, start_row=2)
    
    # Spaltenbreiten anpassen
    apply_column_widths(worksheet, df)
    
    # Freeze Panes (Kopfzeile fixieren)
    if title:
        worksheet.freeze_panes = 'A8'
    else:
        worksheet.freeze_panes = 'A2'


def format_destination_stats(df):
    """Formatiert Destination-Statistiken für bessere Lesbarkeit"""
    df = df.copy()
    
    # Sortiere nach Gesamt_PRM absteigend
    if 'Gesamt_PRM' in df.columns:
        df = df.sort_values('Gesamt_PRM', ascending=False)
    
    # Fülle NaN mit 0
    df = df.fillna(0)
    
    # Integer-Konvertierung für Zählspalten
    int_cols = [col for col in df.columns if col.startswith('Total_') or col in ['<36h', '>36h', 'Gesamt_PRM']]
    for col in int_cols:
        if col in df.columns:
            df[col] = df[col].astype(int)
    
    return df


def format_airline_stats(df):
    """Formatiert Airline-Statistiken für bessere Lesbarkeit"""
    df = df.copy()
    
    # Sortiere nach Gesamt_PRM absteigend
    if 'Gesamt_PRM' in df.columns:
        df = df.sort_values('Gesamt_PRM', ascending=False)
    
    # Fülle NaN mit 0
    df = df.fillna(0)
    
    # Integer-Konvertierung für Zählspalten
    int_cols = [col for col in df.columns if col.startswith('Total_') or col in ['<36h', '>36h', 'Gesamt_PRM']]
    for col in int_cols:
        if col in df.columns:
            df[col] = df[col].astype(int)
    
    return df
